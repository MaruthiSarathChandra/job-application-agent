import html
import re
from pathlib import Path
from typing import Dict, List

from docx import Document
from docx.shared import Inches, Pt


TAG_RE = re.compile(r"<[^>]+>")
WORD_RE = re.compile(r"[a-z0-9+#.]+")


def _tokens(text: str):
    text = html.unescape(TAG_RE.sub(" ", text or "")).lower()
    return set(WORD_RE.findall(text))


def _bullet_text(item):
    if isinstance(item, str):
        return item.strip()
    if isinstance(item, dict):
        return str(item.get("text", "")).strip()
    return ""


def _bullet_tags(item):
    if isinstance(item, dict):
        tags = item.get("tags", []) or []
        if isinstance(tags, str):
            tags = [tags]
        return [str(tag) for tag in tags]
    return []


def _score_bullet(item, job_text: str) -> int:
    job_tokens = _tokens(job_text)
    bullet_tokens = _tokens(_bullet_text(item) + " " + " ".join(_bullet_tags(item)))
    return len(job_tokens.intersection(bullet_tokens))


def _rank_bullets(items, job_text: str, limit: int):
    ranked = []
    for index, item in enumerate(items or []):
        text = _bullet_text(item)
        if not text:
            continue
        ranked.append((_score_bullet(item, job_text), -index, text))
    ranked.sort(reverse=True)
    return [text for _, _, text in ranked[:limit]]


def _set_margins(document: Document):
    section = document.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)


def _set_default_font(document: Document):
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)


def _legacy_location(value):
    if isinstance(value, dict):
        return ", ".join(
            str(value.get(key, "")).strip()
            for key in ("city", "state", "country")
            if str(value.get(key, "")).strip()
        )
    return str(value or "").strip()


def _legacy_dates(item):
    start_month = str(item.get("start_month", "") or "").strip()
    start_year = str(item.get("start_year", "") or "").strip()
    start = " ".join(part for part in (start_month, start_year) if part)

    if item.get("current_job"):
        end = "Present"
    else:
        end_month = str(item.get("end_month", "") or "").strip()
        end_year = str(item.get("end_year", "") or "").strip()
        end = " ".join(part for part in (end_month, end_year) if part)

    return " - ".join(part for part in (start, end) if part)


def _resume_content(profile):
    """
    Return verified resume content, accepting both the V2 resume_content shape and
    the legacy profile shape already used by the Workday automation.

    This adapter does not invent facts. It only converts fields that already
    exist in candidate_profile.yaml.
    """
    content = profile.get("resume_content", {}) or {}
    experience = list(content.get("experience", []) or [])
    projects = list(content.get("projects", []) or [])

    if not experience:
        for item in profile.get("work_experience", []) or []:
            if not isinstance(item, dict):
                continue
            description = str(item.get("description", "") or "").strip()
            bullets = item.get("bullets", []) or []
            if not bullets and description:
                bullets = [description]
            experience.append(
                {
                    "employer": str(item.get("employer", "") or "").strip(),
                    "title": str(
                        item.get("job_title", "")
                        or item.get("title", "")
                        or ""
                    ).strip(),
                    "location": _legacy_location(item.get("location")),
                    "dates": _legacy_dates(item),
                    "bullets": bullets,
                }
            )

    if not projects:
        legacy_projects = profile.get("verified_facts.projects", []) or []
        for item in legacy_projects:
            if isinstance(item, dict):
                name = str(item.get("name", "") or item.get("title", "") or "").strip()
                tech = item.get("tech", "") or item.get("technologies", "") or ""
                if isinstance(tech, list):
                    tech = ", ".join(str(x) for x in tech)
                bullets = item.get("bullets", []) or []
                description = str(item.get("description", "") or "").strip()
                if not bullets and description:
                    bullets = [description]
                projects.append(
                    {
                        "name": name,
                        "tech": str(tech).strip(),
                        "bullets": bullets,
                    }
                )
            elif isinstance(item, str) and item.strip():
                projects.append({"name": item.strip(), "tech": "", "bullets": []})

    return {
        "headline": str(content.get("headline", "") or "").strip(),
        "summary": str(content.get("summary", "") or "").strip(),
        "experience": experience,
        "projects": projects,
    }


def build_tailored_resume(profile, job, output_path: str) -> str:
    """
    Build a DOCX only from explicitly verified local profile content.

    V2 prefers resume_content, but legacy work_experience and
    verified_facts.projects are also accepted so existing profiles do not need a
    manual migration before the first V2 run.
    """
    content = _resume_content(profile)
    experience = content.get("experience", []) or []
    projects = content.get("projects", []) or []

    if not experience and not projects:
        raise ValueError(
            "candidate_profile.yaml needs verified resume_content, work_experience, "
            "or verified_facts.projects before automatic resume generation."
        )

    candidate = profile.get("candidate", {}) or {}
    education = profile.get("education", {}) or {}
    verified_skills = profile.get("verified_facts.skills", []) or []

    job_text = f"{getattr(job, 'title', '')}\n{getattr(job, 'description', '')}"
    job_tokens = _tokens(job_text)

    ranked_skills = []
    for skill in verified_skills:
        skill = str(skill).strip()
        if not skill:
            continue
        score = len(_tokens(skill).intersection(job_tokens))
        ranked_skills.append((score, skill))
    ranked_skills.sort(key=lambda pair: (pair[0], pair[1].lower()), reverse=True)
    skills = [skill for _, skill in ranked_skills]

    document = Document()
    _set_margins(document)
    _set_default_font(document)

    name = str(candidate.get("legal_name") or candidate.get("name") or "").strip()
    if name:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(name)
        run.bold = True
        run.font.size = Pt(15)

    contact_parts = []
    for key in ("email", "phone", "location", "linkedin", "github"):
        value = candidate.get(key)
        if value:
            contact_parts.append(str(value))
    if contact_parts:
        document.add_paragraph(" | ".join(contact_parts))

    headline = str(content.get("headline", "")).strip()
    if not headline and experience:
        headline = str(experience[0].get("title", "") or "").strip()
    if headline:
        p = document.add_paragraph()
        r = p.add_run(headline)
        r.bold = True

    summary = str(content.get("summary", "")).strip()
    if summary:
        document.add_heading("SUMMARY", level=2)
        document.add_paragraph(summary)

    if skills:
        document.add_heading("TECHNICAL SKILLS", level=2)
        document.add_paragraph(" • ".join(skills[:16]))

    if experience:
        document.add_heading("EXPERIENCE", level=2)
        for item in experience:
            employer = str(item.get("employer", "")).strip()
            title = str(item.get("title", "")).strip()
            location = str(item.get("location", "")).strip()
            dates = str(item.get("dates", "")).strip()

            heading = " — ".join(part for part in (title, employer) if part)
            if heading:
                p = document.add_paragraph()
                r = p.add_run(heading)
                r.bold = True
            meta = " | ".join(part for part in (location, dates) if part)
            if meta:
                document.add_paragraph(meta)

            for bullet in _rank_bullets(item.get("bullets", []), job_text, 5):
                document.add_paragraph(bullet, style="List Bullet")

    if projects:
        document.add_heading("PROJECTS", level=2)
        ranked_projects = []
        for index, project in enumerate(projects):
            combined = (
                str(project.get("name", ""))
                + " "
                + str(project.get("tech", ""))
                + " "
                + " ".join(_bullet_text(x) for x in project.get("bullets", []) or [])
            )
            score = len(_tokens(combined).intersection(job_tokens))
            ranked_projects.append((score, -index, project))
        ranked_projects.sort(key=lambda item: (item[0], item[1]), reverse=True)

        for _, _, project in ranked_projects[:4]:
            project_name = str(project.get("name", "")).strip()
            tech = str(project.get("tech", "")).strip()
            if project_name:
                p = document.add_paragraph()
                r = p.add_run(project_name)
                r.bold = True
                if tech:
                    p.add_run(f" | {tech}")
            for bullet in _rank_bullets(project.get("bullets", []), job_text, 3):
                document.add_paragraph(bullet, style="List Bullet")

    if education:
        document.add_heading("EDUCATION", level=2)
        degree = str(education.get("highest_degree", "")).strip()
        university = str(education.get("university", "")).strip()
        graduation = str(education.get("graduation_date", "")).strip()
        line = " | ".join(part for part in (degree, university, graduation) if part)
        if line:
            document.add_paragraph(line)

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return str(path)
